import csv
import json
import tempfile
import uuid
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook


class QuestionnairePreviewStore:
    def __init__(self, root: Path) -> None:
        self.root = root / "questionnaire-previews"

    def save(self, filename: str, data: bytes) -> str:
        token = str(uuid.uuid4())
        destination = self.root / f"{token}-{Path(filename).name}"
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(data)
        return token

    def find(self, token: str) -> Path:
        matches = list(self.root.glob(f"{token}-*"))
        if len(matches) != 1:
            raise HTTPException(status_code=404, detail="Questionnaire preview has expired. Upload the file again.")
        return matches[0]

    def remove(self, token: str) -> None:
        path = self.find(token)
        path.unlink(missing_ok=True)


def read_rows(path: Path) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.DictReader(file)
            columns = list(reader.fieldnames or [])
            rows = [(index, {key: str(value or "") for key, value in row.items()}) for index, row in enumerate(reader, start=2)]
    elif suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        columns = [str(value or "").strip() for value in headers or []]
        rows = [(index, {columns[column]: str(value or "") for column, value in enumerate(row) if column < len(columns)}) for index, row in enumerate(values, start=2)]
        workbook.close()
    else:
        raise HTTPException(status_code=415, detail="Only CSV and XLSX questionnaires are supported.")
    if not columns or not any(columns):
        raise HTTPException(status_code=400, detail="The questionnaire has no usable header row.")
    return columns, rows


def inspect_questionnaire(path: Path, token: str) -> dict:
    columns, rows = read_rows(path)
    return {"token": token, "filename": path.name.removeprefix(f"{token}-"), "columns": columns, "row_count": len(rows), "sample_rows": [row for _, row in rows[:8]]}
